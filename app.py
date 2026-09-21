#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
排班表转ICS日历工具
输入固定格式的排班Excel，输出每人独立的ics日历文件。
"""

import os
import sys
import traceback
from datetime import datetime, timedelta

import tkinter as tk
from tkinter import ttk, filedialog, messagebox, scrolledtext

try:
    import openpyxl
except ImportError:
    openpyxl = None

try:
    from pypinyin import lazy_pinyin
except ImportError:
    lazy_pinyin = None


# ── 班次时间映射（固定格式班表） ──────────────────────────────
SHIFT_TIMES = {
    '早班': ('09:30', '17:30'),
    '早值': ('09:30', '20:00'),
    '中班': ('13:00', '20:00'),
    '晚班': ('15:00', '22:00'),
}
REST_KEYWORD = '休息'


def name_to_pinyin(name: str) -> str:
    """中文名转拼音（无调号，小写连写）。"""
    if lazy_pinyin:
        return ''.join(lazy_pinyin(name))
    # 兜底：已知姓名映射
    fallback = {
        '朱昌': 'zhuchang', '刘颂庆': 'liusongqing', '王铖': 'wangcheng',
        '李丽丽': 'lilili', '林胜捷': 'linshengjie', '农汉翔': 'nonghanxiang',
    }
    return fallback.get(name, name)


def parse_schedule(xlsx_path: str, year: int, month: int):
    """
    解析固定格式排班表。
    返回: (dates, people_list)
    dates = [(col_index, date_obj), ...]
    people_list = [{'name': ..., 'days': [(date_obj, shift_str), ...]}, ...]

    跨月处理：第1行日期只含「日」。当日期相对前一列发生回退
    （如 31 后出现 1），或超出当月最大天数时，自动进位到下个月；
    12 月进位后自动跨到下一年。
    """
    import calendar

    if not xlsx_path.lower().endswith('.xlsx'):
        raise ValueError(
            '仅支持 .xlsx 格式的排班表。\n'
            '如果手头是旧版 .xls 文件，请先用 Excel 或 WPS「另存为」.xlsx 格式，再重新选择。')

    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb.active

    # 第1行：A1='日期'，B1起为日期（日）
    # 第2行：A2='姓名'，B2起为星期
    # 第3行起：A列为姓名，B-H列为班次
    day_cols = []  # [(col_index, day_int)]
    for col in range(2, ws.max_column + 1):
        val = ws.cell(row=1, column=col).value
        if val is None:
            break
        try:
            day_int = int(val)
            day_cols.append((col, day_int))
        except (ValueError, TypeError):
            break

    if not day_cols:
        raise ValueError('未在第1行检测到日期数据，请确认班表格式。')

    # 根据初始年月逐列确定真实日期：遇回退或超当月天数 → 自动进下月
    dates = []  # [(col_index, date_obj)]
    cur_year, cur_month = year, month
    prev_day = None
    for col, day_int in day_cols:
        max_day = calendar.monthrange(cur_year, cur_month)[1]
        if prev_day is not None and (day_int < prev_day or day_int > max_day):
            # 跨月：进入下个月
            cur_month += 1
            if cur_month > 12:
                cur_month = 1
                cur_year += 1
            max_day = calendar.monthrange(cur_year, cur_month)[1]
        if day_int > max_day:
            raise ValueError(
                f'日期 {day_int} 超出 {cur_year}年{cur_month}月最大天数 {max_day} 天，请检查班表。')
        dates.append((col, datetime(cur_year, cur_month, day_int)))
        prev_day = day_int

    # 人员行：从第3行开始，A列有姓名
    people = []
    for row in range(3, ws.max_row + 1):
        name = ws.cell(row=row, column=1).value
        if name is None or str(name).strip() == '':
            break
        # 跳过底部说明行（包含冒号的说明文字）
        name_str = str(name).strip()
        if '：' in name_str or ':' in name_str:
            continue
        schedule = []
        for col, date_obj in dates:
            shift = ws.cell(row=row, column=col).value
            if shift is not None and str(shift).strip():
                schedule.append((date_obj, str(shift).strip()))
        people.append({'name': name_str, 'days': schedule})

    if not people:
        raise ValueError('未检测到人员排班数据，请确认班表格式。')

    return dates, people


def generate_ics(name: str, days_schedule) -> str:
    """为单人生成ICS内容字符串。days_schedule = [(date_obj, shift_str), ...]"""
    pinyin = name_to_pinyin(name)
    events = []
    dtstamp = datetime.now().strftime('%Y%m%dT%H%M%S')

    for idx, (date_obj, shift) in enumerate(days_schedule, start=1):
        uid = f'{pinyin}-{date_obj.strftime("%Y%m%d")}-{idx}@schedule'
        summary = f'{name} - {shift}'

        if shift == REST_KEYWORD:
            dtstart = date_obj.strftime('%Y%m%d')
            dtend = (date_obj + timedelta(days=1)).strftime('%Y%m%d')
            events.append(
                f'BEGIN:VEVENT\n'
                f'UID:{uid}\n'
                f'DTSTAMP:{dtstamp}\n'
                f'DTSTART;VALUE=DATE:{dtstart}\n'
                f'DTEND;VALUE=DATE:{dtend}\n'
                f'SUMMARY:{summary}\n'
                f'END:VEVENT'
            )
        elif shift in SHIFT_TIMES:
            start_t, end_t = SHIFT_TIMES[shift]
            sh, sm = map(int, start_t.split(':'))
            eh, em = map(int, end_t.split(':'))
            dtstart = datetime(date_obj.year, date_obj.month, date_obj.day, sh, sm)
            dtend = datetime(date_obj.year, date_obj.month, date_obj.day, eh, em)
            events.append(
                f'BEGIN:VEVENT\n'
                f'UID:{uid}\n'
                f'DTSTAMP:{dtstamp}\n'
                f'DTSTART;TZID=Asia/Shanghai:{dtstart.strftime("%Y%m%dT%H%M%S")}\n'
                f'DTEND;TZID=Asia/Shanghai:{dtend.strftime("%Y%m%dT%H%M%S")}\n'
                f'SUMMARY:{summary}\n'
                f'END:VEVENT'
            )
        # 未知班次跳过

    ics = (
        'BEGIN:VCALENDAR\n'
        'VERSION:2.0\n'
        f'PRODID:-//Schedule//{name}//CN\n'
        'CALSCALE:GREGORIAN\n'
        'METHOD:PUBLISH\n'
        f'X-WR-CALNAME:{name}的排班\n'
        'X-WR-TIMEZONE:Asia/Shanghai\n'
        'BEGIN:VTIMEZONE\n'
        'TZID:Asia/Shanghai\n'
        'X-LIC-LOCATION:Asia/Shanghai\n'
        'BEGIN:STANDARD\n'
        'TZOFFSETFROM:+0800\n'
        'TZOFFSETTO:+0800\n'
        'TZNAME:CST\n'
        'DTSTART:19700101T000000\n'
        'END:STANDARD\n'
        'END:VTIMEZONE\n'
        + '\n'.join(events) + '\n'
        'END:VCALENDAR'
    )
    return ics


def process(xlsx_path: str, year: int, month: int, log_fn):
    """主处理流程，返回输出目录。"""
    if openpyxl is None:
        raise RuntimeError('缺少依赖 openpyxl，请先运行 pip install openpyxl')

    log_fn(f'读取文件: {xlsx_path}')
    dates, people = parse_schedule(xlsx_path, year, month)
    log_fn(f'检测到 {len(dates)} 天、{len(people)} 人')
    log_fn(f'排班日期: {dates[0][1].strftime("%Y-%m-%d")} 至 {dates[-1][1].strftime("%Y-%m-%d")}')

    out_dir = os.path.join(os.path.dirname(os.path.abspath(xlsx_path)), 'ics_output')
    os.makedirs(out_dir, exist_ok=True)

    for person in people:
        ics_content = generate_ics(person['name'], person['days'])
        pinyin = name_to_pinyin(person['name'])
        out_path = os.path.join(out_dir, f'{pinyin}.ics')
        with open(out_path, 'w', encoding='utf-8') as f:
            f.write(ics_content)
        log_fn(f'  生成: {pinyin}.ics ({len(person["days"])} 个日程) - {person["name"]}')

    log_fn(f'完成！输出目录: {out_dir}')
    return out_dir


# ── GUI ────────────────────────────────────────────────────────
class App:
    def __init__(self, root):
        self.root = root
        self.root.title('排班表转ICS日历')
        self.root.geometry('620x480')
        self.root.minsize(560, 420)

        self.file_path = tk.StringVar()
        now = datetime.now()
        self.year_var = tk.StringVar(value=str(now.year))
        self.month_var = tk.StringVar(value=str(now.month))

        self._build_ui()

    def _build_ui(self):
        pad = {'padx': 12, 'pady': 6}

        # 顶部：年月选择
        top = ttk.Frame(self.root)
        top.pack(fill='x', **pad)

        ttk.Label(top, text='排班年份:').pack(side='left')
        year_cb = ttk.Combobox(top, textvariable=self.year_var, width=6,
                                values=[str(y) for y in range(2020, 2036)])
        year_cb.pack(side='left', padx=(4, 16))

        ttk.Label(top, text='排班月份:').pack(side='left')
        month_cb = ttk.Combobox(top, textvariable=self.month_var, width=4,
                                 values=[str(m) for m in range(1, 13)])
        month_cb.pack(side='left', padx=(4, 16))

        ttk.Label(top, text='（班表仅含日期，需指定起始年月；跨月自动进位）',
                  foreground='gray').pack(side='left')

        # 文件选择
        file_frame = ttk.Frame(self.root)
        file_frame.pack(fill='x', **pad)

        ttk.Button(file_frame, text='选择班表Excel', command=self._choose_file).pack(side='left')
        ttk.Entry(file_frame, textvariable=self.file_path, state='readonly').pack(
            side='left', fill='x', expand=True, padx=8)

        # 生成按钮
        btn_frame = ttk.Frame(self.root)
        btn_frame.pack(fill='x', **pad)
        self.gen_btn = ttk.Button(btn_frame, text='生成ICS日历', command=self._generate)
        self.gen_btn.pack(side='left')

        # 日志
        log_frame = ttk.LabelFrame(self.root, text='运行日志')
        log_frame.pack(fill='both', expand=True, padx=12, pady=(0, 12))
        self.log = scrolledtext.ScrolledText(log_frame, height=12, state='disabled',
                                              font=('Menlo', 11) if sys.platform == 'darwin' else ('Consolas', 10))
        self.log.pack(fill='both', expand=True, padx=6, pady=6)

        self._log('请选择排班表Excel文件，确认年月后点击生成。')

    def _log(self, msg: str):
        self.log.configure(state='normal')
        self.log.insert('end', msg + '\n')
        self.log.see('end')
        self.log.configure(state='disabled')
        self.root.update_idletasks()

    def _choose_file(self):
        path = filedialog.askopenfilename(
            title='选择排班表Excel',
            filetypes=[('Excel文件', '*.xlsx *.xls'), ('所有文件', '*.*')]
        )
        if path:
            self.file_path.set(path)
            self._log(f'已选择: {path}')

    def _generate(self):
        path = self.file_path.get().strip()
        if not path:
            messagebox.showwarning('提示', '请先选择排班表Excel文件。')
            return
        try:
            year = int(self.year_var.get())
            month = int(self.month_var.get())
        except ValueError:
            messagebox.showwarning('提示', '年份或月份格式不正确。')
            return

        self.gen_btn.configure(state='disabled')
        try:
            out_dir = process(path, year, month, self._log)
            messagebox.showinfo('完成', f'ICS文件已生成到:\n{out_dir}')
            # macOS 上打开输出目录
            if sys.platform == 'darwin':
                os.system(f'open "{out_dir}"')
        except Exception as e:
            self._log(f'错误: {e}')
            self._log(traceback.format_exc())
            messagebox.showerror('出错了', str(e))
        finally:
            self.gen_btn.configure(state='normal')


def main():
    root = tk.Tk()
    # macOS 风格优化
    if sys.platform == 'darwin':
        try:
            from tkinter import font as tkfont
            default_font = tkfont.nametofont('TkDefaultFont')
            default_font.configure(size=13)
        except Exception:
            pass
    App(root)
    root.mainloop()


if __name__ == '__main__':
    main()
